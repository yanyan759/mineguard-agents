"""真实调用配置 API，捕获上传未消费、版本漂移、跨矿配置和示例规则泄漏。"""
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from app.api.server import app


def settings():
    return dict(mine_id="MINE-ACCEPT", roadway_id="RDW-01", area="工作面-1", version="v1",
                source="验收专用参数文件", valid_from="2026-01-01T00:00:00Z", valid_to="2027-01-01T00:00:00Z",
                window_minutes=60, rules=[dict(metric="max_energy_j", operator=">=", threshold=5000, level="yellow")],
                assets=[], inventory_source="input", max_inventory_age_hours=24)


def task():
    return dict(input_type="raw_events", mine_id="MINE-ACCEPT", roadway_id="RDW-01", area="工作面-1",
                source="验收记录", method_id="statistics-v1", window_start="2026-09-16T00:00:00Z",
                window_end="2026-09-16T01:00:00Z", events=[dict(event_id="a",timestamp="2026-09-16T00:30:00Z",energy=6,energy_unit="kJ")])


class ActualPipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.env = patch.dict(os.environ, {"MINE_CONFIGURATION_DB":str(Path(self.tmp.name)/"config.sqlite3")})
        self.env.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close(); self.env.stop(); self.tmp.cleanup()

    def upload(self, category, content, filename="asset.json"):
        meta = dict(mine_id="MINE-ACCEPT",roadway_id="RDW-01",area="工作面-1",category=category,
                    version="v1",source="验收资料",title=filename,observed_at="2026-09-16T00:00:00Z")
        return self.client.post("/api/configuration/ingest",data={"metadata":json.dumps(meta)},
                                files={"file":(filename,content.encode("utf-8"),"application/octet-stream")})

    def prepare(self, value):
        return self.client.post("/api/configuration/prepare",json={"query":json.dumps(value,ensure_ascii=False)})

    def test_uploaded_inventory_is_consumed_and_configuration_snapshot_is_immutable(self):
        asset=self.upload("resources","group,resource_id,name,available,required,unit\npersonnel,p,支护工,8,6,人\nequipment,e,钻机,1,2,台\n","inventory.csv")
        self.assertEqual(asset.status_code,200,asset.text)
        config=settings();config["assets"]=[asset.json()["asset_id"]]
        saved=self.client.post("/api/configuration",json=config)
        self.assertEqual(saved.status_code,200,saved.text)
        prepared=self.prepare(task());self.assertEqual(prepared.status_code,200,prepared.text)
        snapshot=prepared.json()["snapshot"]
        self.assertEqual(snapshot["task"]["resources"]["equipment"][0]["available"],1)
        self.assertEqual(snapshot["assessment"]["risk_level"],"yellow")
        config["version"]="v2";config["rules"][0]["threshold"]=7000
        self.assertEqual(self.client.post("/api/configuration",json=config).status_code,200)
        self.assertEqual(self.prepare(task()).json()["snapshot"]["assessment"]["risk_level"],"normal")
        self.assertEqual(snapshot["assessment"]["risk_level"],"yellow")
        self.assertEqual(snapshot["configuration"]["version"],"v1")

    def test_cross_mine_assets_and_duplicate_versions_are_rejected(self):
        asset=self.upload("case",json.dumps({"text":"微震与顶板异常调查记录"},ensure_ascii=False))
        self.assertEqual(asset.status_code,200,asset.text)
        config=settings();config["mine_id"]="OTHER";config["assets"]=[asset.json()["asset_id"]]
        self.assertEqual(self.client.post("/api/configuration",json=config).status_code,422)
        config=settings();self.assertEqual(self.client.post("/api/configuration",json=config).status_code,200)
        config["rules"][0]["threshold"]=9
        self.assertEqual(self.client.post("/api/configuration",json=config).status_code,409)

    def test_invalid_inventory_does_not_become_indexed(self):
        result=self.upload("resources",'{"equipment":[{"resource_id":"e","name":"钻机","available":-1,"required":2,"unit":"台"}]}')
        self.assertEqual(result.status_code,422,result.text)

    def test_word_and_pdf_text_are_parsed_and_cited(self):
        from io import BytesIO
        from docx import Document
        from reportlab.pdfgen.canvas import Canvas
        meta=dict(mine_id="MINE-ACCEPT",roadway_id="RDW-01",area="工作面-1",category="regulation",version="v1",source="验收文件",title="文件解析验收",observed_at="2026-09-16T00:00:00Z")
        word=Document();word.add_paragraph("Document paragraph evidence")
        word.add_table(rows=1,cols=1).cell(0,0).text="Table evidence"
        word_buffer=BytesIO();word.save(word_buffer)
        pdf_buffer=BytesIO();pdf=Canvas(pdf_buffer);pdf.drawString(50,700,"PDF text evidence");pdf.save()
        for filename,raw,expected in [("policy.docx",word_buffer.getvalue(),"Table evidence"),("policy.pdf",pdf_buffer.getvalue(),"PDF text evidence")]:
            response=self.client.post("/api/configuration/ingest",data={"metadata":json.dumps(meta)},files={"file":(filename,raw)})
            self.assertEqual(response.status_code,200,response.text)
            self.assertIn(expected,response.json()["parsed"]["text"])
        broken=self.client.post("/api/configuration/ingest",data={"metadata":json.dumps(meta)},files={"file":("bad.pdf",b"not a pdf")})
        self.assertEqual(broken.status_code,422)

    def test_plan_unit_mismatch_and_stale_inventory_fail_before_start(self):
        inventory=self.upload("resources",'{"equipment":[{"resource_id":"e","name":"钻机","required":1,"available":2,"unit":"台"}]}').json()
        plan=self.upload("support",json.dumps({"text":"验收方案","actions":[{"sequence":1,"action":"检查","owner_role":"检查人员","deadline_minutes":None}],"requirements":[{"group":"equipment","name":"钻机","required":1,"unit":"组"}]},ensure_ascii=False)).json()
        config=settings();config.update(assets=[inventory["asset_id"],plan["asset_id"]],plan_asset_id=plan["asset_id"])
        self.assertEqual(self.client.post("/api/configuration",json=config).status_code,200)
        self.assertEqual(self.prepare(task()).status_code,422)
        config.update(version="v2",plan_asset_id=None,max_inventory_age_hours=.1)
        self.client.post("/api/configuration",json=config)
        self.assertEqual(self.prepare(task()).status_code,422)

    def test_real_summary_without_configuration_never_uses_simulation_thresholds(self):
        result=self.client.post("/api/configuration/prepare",json={"query":"MINE-021 / RDW-031 / 工作面-31：过去45分钟监测到8条微震事件，频次由7.2次/h升至14.6次/h，最大能量62000 J，b值由0.94降至0.71，数据缺失率2%。高能事件2条。"})
        self.assertEqual(result.status_code,200,result.text)
        assessment=result.json()["snapshot"]["assessment"]
        self.assertIsNone(assessment["risk_level"])
        self.assertEqual(assessment["metrics"]["max_energy_j"],62000)
        self.assertEqual(assessment["method_id"],"observations-only")

    def test_expired_or_window_mismatched_parameters_fail_preflight(self):
        config=settings();config["valid_to"]="2026-09-01T00:00:00Z"
        self.assertEqual(self.client.post("/api/configuration",json=config).status_code,200)
        result=self.prepare(task());self.assertEqual(result.status_code,422,result.text)

    def test_text_entry_uses_arbitrary_configured_scope_time_and_version(self):
        config=settings();config.update(mine_id="ACCEPT-SITE",roadway_id="LANE-A",area="东区测点")
        self.assertEqual(self.client.post("/api/configuration",json=config).status_code,200)
        query="ACCEPT-SITE / LANE-A / 东区测点：过去60分钟监测到3条微震事件，频次由2次/h升至3次/h，最大能量6000 J，b值由1.4降至1.2，缺失率0%。配置版本：v1；观测截止：2026-09-16T01:00:00Z；数据来源：现场监测摘要。"
        result=self.client.post("/api/configuration/prepare",json={"query":query})
        self.assertEqual(result.status_code,200,result.text)
        self.assertEqual(result.json()["snapshot"]["assessment"]["risk_level"],"yellow")
        self.assertEqual(result.json()["task"]["source"],"现场监测摘要")
        self.assertEqual(result.json()["task"]["configuration_version"],"v1")

    def test_six_nodes_consume_frozen_documents_plan_and_quantities(self):
        from app.agent.workflow import build_workflow
        from app.agent.workflow_state import new_workflow_state
        from app.schemas.workflow import WorkflowContext
        from datetime import datetime, timezone
        inventory=self.upload("resources",'group,resource_id,name,available,required,unit\nequipment,e,钻机,1,2,台\n',"inventory.csv").json()
        case=self.upload("case",json.dumps({"text":"本测区微震观测案例原文"},ensure_ascii=False)).json()
        regulation=self.upload("regulation",json.dumps({"text":"本矿批准规程原文"},ensure_ascii=False)).json()
        plan=self.upload("support",json.dumps({"text":"验收方案", "actions":[{"sequence":1,"action":"核对钻机状态并复查测点","owner_role":"支护负责人","deadline_minutes":20}],"requirements":[{"group":"equipment","name":"钻机","required":2,"unit":"台"}]},ensure_ascii=False)).json()
        config=settings();config.update(assets=[a["asset_id"] for a in [inventory,case,regulation,plan]],plan_asset_id=plan["asset_id"])
        self.assertEqual(self.client.post("/api/configuration",json=config).status_code,200)
        frozen=self.prepare(task()).json()["snapshot"]
        config["version"]="v2";config["rules"][0]["threshold"]=7000
        self.client.post("/api/configuration",json=config)
        graph=build_workflow().compile()
        def run(snapshot):
            state=new_workflow_state(WorkflowContext(workflow_run_id="acceptance",trace_id="acceptance",thread_id="acceptance",mine_id="MINE-ACCEPT",roadway_id="RDW-01",started_at=datetime.now(timezone.utc)))
            state.update(custom_task=snapshot["task"],actual_snapshot=snapshot)
            return graph.invoke(state)
        result=run(frozen)
        self.assertEqual(result["perception_result"].risk_level.value,"yellow")
        self.assertEqual(result["knowledge_result"].similar_cases[0].case_id,case["asset_id"])
        self.assertEqual(result["knowledge_result"].regulations[0].requirement,"本矿批准规程原文")
        self.assertEqual(result["decision_result"].primary_plan.actions[0].action,"核对钻机状态并复查测点")
        self.assertEqual(result["resource_result"].resource_gaps[0].available,1)
        self.assertIsNotNone(result["control_result"])
        self.assertIsNotNone(result["reflection_result"])
        actual=task();actual["resources"]={"equipment":[dict(resource_id="e",name="钻机",available=3,required=0,unit="台")]}
        result=run(self.prepare(actual).json()["snapshot"])
        self.assertEqual(result["perception_result"].risk_level.value,"normal")
        self.assertEqual(result["resource_result"].quantity_check,"sufficient")
        self.assertIsNotNone(result["reflection_result"])


if __name__ == "__main__": unittest.main()
