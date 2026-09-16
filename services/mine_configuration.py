"""按矿井保存不可变配置、解析资料并冻结本次运行输入。"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import sqlite3
from contextlib import contextmanager, closing
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Literal

from pydantic import Field, model_validator
from app.schemas.workflow import ResourceSnapshot, DispatchAction
from app.services.raw_monitoring import StrictModel, RawMonitoringTask, analyze_raw

LEVELS = ["normal", "blue", "yellow", "orange", "red"]
GROUPS = ("personnel", "support_materials", "equipment")
METRICS = ("max_energy_j", "event_count", "frequency_per_hour", "b_value", "missing_rate",
           "frequency_increase", "b_drop", "high_energy_count", "spatial_cluster_fraction")


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)


def digest(value):
    return hashlib.sha256(canonical(value).encode()).hexdigest()


class Scope(StrictModel):
    mine_id: str = Field(min_length=1, max_length=100, pattern=r"^[\w-]+$")
    roadway_id: str = Field(min_length=1, max_length=100, pattern=r"^[\w-]+$")
    area: str = Field(min_length=1, max_length=200)


def scope_key(value):
    return canonical({k: value[k] for k in ("mine_id", "roadway_id", "area")})


class Rule(StrictModel):
    metric: Literal["max_energy_j", "event_count", "frequency_per_hour", "b_value", "missing_rate",
                    "frequency_increase", "b_drop", "high_energy_count", "spatial_cluster_fraction"]
    operator: Literal[">=", "<="]
    threshold: float = Field(ge=0)
    level: Literal["blue", "yellow", "orange", "red"]


class MineConfiguration(Scope):
    version: str = Field(min_length=1, max_length=100)
    source: str = Field(min_length=1, max_length=1000)
    valid_from: datetime
    valid_to: datetime
    window_minutes: float = Field(gt=0, le=525600)
    rules: list[Rule] = Field(default_factory=list, max_length=50)
    assets: list[str] = Field(default_factory=list, max_length=100)
    inventory_source: Literal["input", "business", "mysql"] = "input"
    max_inventory_age_hours: float = Field(default=24, gt=0, le=8760)
    recheck_minutes: int | None = Field(default=None, gt=0)
    approval_reference: str = ""
    plan_asset_id: str | None = None

    @model_validator(mode="after")
    def coherent(self):
        if any(x.tzinfo is None for x in (self.valid_from, self.valid_to)) or self.valid_to <= self.valid_from:
            raise ValueError("配置有效期必须包含时区且结束晚于开始")
        if len(set(self.assets)) != len(self.assets):
            raise ValueError("资料编号不能重复")
        if self.plan_asset_id and self.plan_asset_id not in self.assets:
            raise ValueError("选定方案必须包含在本配置资料清单中")
        return self


class AssetMetadata(Scope):
    category: Literal["resources", "case", "regulation", "support"]
    title: str = Field(min_length=1, max_length=300)
    version: str = Field(min_length=1, max_length=100)
    source: str = Field(min_length=1, max_length=2000)
    observed_at: datetime
    valid_from: datetime | None = None
    valid_to: datetime | None = None

    @model_validator(mode="after")
    def dates(self):
        if self.observed_at.tzinfo is None:
            raise ValueError("资料时间需要时区")
        if (self.valid_from is None) != (self.valid_to is None):
            raise ValueError("资料适用期需要同时填写起止时间")
        if self.valid_from and (self.valid_from.tzinfo is None or self.valid_to.tzinfo is None or self.valid_to <= self.valid_from):
            raise ValueError("资料适用期不合法")
        return self


class Requirement(StrictModel):
    group: Literal["personnel", "support_materials", "equipment"]
    name: str = Field(min_length=1)
    unit: str = Field(min_length=1)
    required: float = Field(gt=0)


class Document(StrictModel):
    text: str = Field(min_length=1, max_length=500000)
    terms: list[str] = Field(default_factory=list, max_length=100)
    actions: list[DispatchAction] = Field(default_factory=list, max_length=100)
    requirements: list[Requirement] = Field(default_factory=list, max_length=200)


def validate_resources(value):
    """资源按名称和单位精确匹配，拒绝重复、非有限值和空清单。"""
    snapshot = ResourceSnapshot.model_validate(value)
    count = 0
    for group in GROUPS:
        names, ids = set(), set()
        for item in getattr(snapshot, group):
            count += 1
            if not item.name.strip() or not item.unit.strip() or not item.resource_id.strip():
                raise ValueError("资源名称、编号和单位不能为空")
            if item.name in names or item.resource_id in ids:
                raise ValueError("同类资源名称或编号重复")
            if any(not math.isfinite(v) for v in (item.available, item.required)):
                raise ValueError("资源数量必须为有限值")
            names.add(item.name); ids.add(item.resource_id)
    if not count:
        raise ValueError("资源台账不能为空")
    return snapshot.model_dump(mode="json")


def parse_asset(filename, raw, meta):
    """只解析支持格式；扫描件无正文时明确拒绝，不猜测数量或施工动作。"""
    suffix = Path(filename).suffix.lower()
    if len(raw) > 20 * 1024 * 1024:
        raise ValueError("单份资料不能超过20MB")
    if meta.category == "resources":
        if suffix == ".json":
            result = json.loads(raw.decode("utf-8-sig"))
        elif suffix == ".csv":
            reader = csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
            columns = ["group", "resource_id", "name", "available", "required", "unit"]
            if not reader.fieldnames or len(reader.fieldnames) != len(set(reader.fieldnames)) or set(reader.fieldnames) != set(columns):
                raise ValueError("资源CSV列应为group,resource_id,name,available,required,unit")
            result = {g: [] for g in GROUPS}
            for row in reader:
                group = row.pop("group")
                if group not in GROUPS or None in row:
                    raise ValueError("资源类别或列数不正确")
                result[group].append(row)
        else:
            raise ValueError("资源台账支持JSON或CSV，请先将文档中的数量整理为模板字段")
        return {"resources": validate_resources(result)}
    if suffix == ".json":
        document = Document.model_validate_json(raw)
    else:
        if suffix in (".txt", ".md"):
            text = raw.decode("utf-8-sig")
        elif suffix == ".pdf":
            from pypdf import PdfReader
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
        elif suffix == ".docx":
            from docx import Document as WordDocument
            word = WordDocument(io.BytesIO(raw))
            text = "\n".join([p.text for p in word.paragraphs] +
                             [" | ".join(c.text for c in row.cells) for table in word.tables for row in table.rows])
        else:
            raise ValueError("知识资料支持JSON、TXT、Markdown、可提取文本的PDF及DOCX")
        if not text.strip():
            raise ValueError("文件无可提取正文，请提供文字版或先完成OCR校对")
        document = Document(text=text.strip())
    if document.actions and meta.category != "support":
        raise ValueError("结构化处置动作只能来自支护方案")
    if len({a.sequence for a in document.actions}) != len(document.actions):
        raise ValueError("方案动作序号重复")
    if len({(r.group, r.name) for r in document.requirements}) != len(document.requirements):
        raise ValueError("方案资源需求重复")
    return document.model_dump(mode="json")


class ConfigurationStore:
    def __init__(self, path=None):
        self.path = Path(path or os.getenv("MINE_CONFIGURATION_DB") or Path(__file__).resolve().parents[1]/"output/mine_configuration.sqlite3")

    @contextmanager
    def connection(self):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with closing(sqlite3.connect(self.path, timeout=20)) as conn, conn:
            conn.row_factory = sqlite3.Row
            conn.executescript("""
                CREATE TABLE IF NOT EXISTS assets(id TEXT PRIMARY KEY, scope TEXT, body TEXT, original BLOB);
                CREATE TABLE IF NOT EXISTS configurations(id TEXT PRIMARY KEY, scope TEXT, version TEXT, body TEXT,
                    created_at TEXT, UNIQUE(scope,version));
                CREATE TABLE IF NOT EXISTS usage(run_id TEXT, asset_id TEXT, PRIMARY KEY(run_id,asset_id));
            """)
            yield conn

    def ingest(self, filename, raw, meta):
        parsed = parse_asset(filename, raw, meta)
        body = {**meta.model_dump(mode="json"), "filename":Path(filename).name,
                "sha256":hashlib.sha256(raw).hexdigest(), "parsed":parsed, "status":"indexed",
                "stages":["uploaded","validated","indexed"]}
        body["asset_id"] = digest(body)
        with self.connection() as conn:
            conn.execute("INSERT OR IGNORE INTO assets VALUES (?,?,?,?)",
                         (body["asset_id"],scope_key(body),canonical(body),raw))
        return body

    def asset(self, asset_id):
        with self.connection() as conn:
            row = conn.execute("SELECT body FROM assets WHERE id=?",(asset_id,)).fetchone()
        if not row:
            raise ValueError("资料不存在，请重新选择已入库文件")
        return json.loads(row[0])

    def save(self, config):
        body = config.model_dump(mode="json")
        resources = 0
        for asset_id in config.assets:
            asset = self.asset(asset_id)
            if scope_key(asset) != scope_key(body):
                raise ValueError("资料矿井/巷道/测区与配置不一致")
            resources += asset["category"] == "resources"
        if resources > 1:
            raise ValueError("同一配置只能选一份库存台账，避免同名资源重复累计")
        if config.plan_asset_id:
            asset = self.asset(config.plan_asset_id)
            if asset["category"] != "support" or not asset["parsed"].get("actions") or not asset["parsed"].get("requirements"):
                raise ValueError("执行方案必须是含actions和requirements的结构化支护资料")
        body["revision"] = digest(body)
        now = datetime.now(timezone.utc).isoformat()
        with self.connection() as conn:
            old = conn.execute("SELECT body FROM configurations WHERE scope=? AND version=?",(scope_key(body),config.version)).fetchone()
            if old:
                if json.loads(old[0]) != body:
                    raise FileExistsError("该版本已保存，修改内容请使用新版本号")
                return body
            conn.execute("INSERT INTO configurations VALUES (?,?,?,?,?)",(body["revision"],scope_key(body),config.version,canonical(body),now))
        return body

    def current(self, scope, version=None):
        with self.connection() as conn:
            if version:
                row = conn.execute("SELECT body FROM configurations WHERE scope=? AND version=?",(scope_key(scope),version)).fetchone()
            else:
                row = conn.execute("SELECT body FROM configurations WHERE scope=? ORDER BY created_at DESC,rowid DESC LIMIT 1",(scope_key(scope),)).fetchone()
        if version and not row:
            raise ValueError("指定的矿井配置版本不存在")
        return json.loads(row[0]) if row else None

    def listing(self, scope):
        with self.connection() as conn:
            rows = conn.execute("SELECT body FROM assets WHERE scope=? ORDER BY rowid DESC",(scope_key(scope),)).fetchall()
            usage = {r[0]:r[1] for r in conn.execute("SELECT asset_id,group_concat(run_id) FROM usage GROUP BY asset_id")}
        assets = []
        for row in rows:
            asset = json.loads(row[0])
            asset["used_in_runs"] = usage.get(asset["asset_id"],"").split(",") if asset["asset_id"] in usage else []
            assets.append(asset)
        return {"configuration":self.current(scope),"assets":assets}

    def mark_used(self, run_id, assets):
        with self.connection() as conn:
            conn.executemany("INSERT OR IGNORE INTO usage VALUES (?,?)",[(run_id,a["asset_id"]) for a in assets])


def task_metrics(task):
    data = task.model_dump(mode="json")
    if data.get("input_type") == "raw_events":
        # statistics-v1只负责观测统计；方法参数随后由本次配置或明确研究方法确定。
        raw = analyze_raw(task.model_copy(update={"method_id":"statistics-v1"}))
        return raw["metrics"], task.window_start, task.window_end, "observed"
    end = getattr(task,"window_end",None) or datetime.now(timezone.utc)
    metrics = {"event_count":task.event_count,"max_energy_j":task.max_energy_j,
               "frequency_per_hour":task.frequency_end,"b_value":task.b_end,"missing_rate":task.missing_rate,
               "window_minutes":task.window_minutes,"high_energy_count":task.high_energy_count,
               "spatial_cluster_fraction":task.spatial_cluster_fraction,
               "frequency_increase":task.frequency_end-task.frequency_start if task.frequency_start is not None else None,
               "b_drop":task.b_start-task.b_end if task.b_start is not None else None}
    return metrics,end-timedelta(minutes=task.window_minutes),end,"observed" if getattr(task,"window_end",None) else "submission_time"


def prepare_actual(task, store=None):
    """API、对话、原始文件共用。返回可序列化快照，节点不再读取可变配置。"""
    store = store or ConfigurationStore()
    data = task.model_dump(mode="json")
    config = store.current(data, data.get("configuration_version"))
    metrics,start,end,time_source = task_metrics(task)
    if data.get("input_type") != "raw_events":
        data["window_end"] = end.isoformat()  # 摘要归档也固定截止时间，重放不重新取当前时钟
    documents = []
    warnings = []
    inventory_asset_id = None
    if config:
        if start < datetime.fromisoformat(config["valid_from"]) or end > datetime.fromisoformat(config["valid_to"]):
            raise ValueError("观测窗口不在配置有效期内，请选择适用参数版本")
        if not math.isclose(metrics["window_minutes"],config["window_minutes"]):
            raise ValueError("观测窗口长度与参数配置不匹配")
        documents = [store.asset(x) for x in config["assets"]]
        for doc in documents:
            if doc.get("valid_from") and (start < datetime.fromisoformat(doc["valid_from"]) or end > datetime.fromisoformat(doc["valid_to"])):
                raise ValueError("资料适用期与观测不匹配："+doc["title"])
        inventory = next((x for x in documents if x["category"]=="resources"),None)
        if inventory and not data.get("resources"):
            age = (end-datetime.fromisoformat(inventory["observed_at"])).total_seconds()/3600
            if age < 0 or age > config["max_inventory_age_hours"]:
                raise ValueError("库存台账时间晚于观测或超过允许时效")
            data["resources"] = inventory["parsed"]["resources"]
            inventory_asset_id = inventory["asset_id"]
        data["configuration_version"] = config["version"]
        data["inventory_source"] = data.get("inventory_source") or config["inventory_source"]
        if data["inventory_source"] == "input" and config["inventory_source"] != "input":
            data["inventory_source"] = config["inventory_source"]
    if data.get("resources"):
        data["resources"] = validate_resources(data["resources"])
    plan = next((x for x in documents if config and x["asset_id"] == config["plan_asset_id"]),None)
    inventory_evidence = None
    if data.get("inventory_source", "input") != "input":
        # 在准入时读取一次业务库存，节点及重启恢复均消费同一份快照。
        if data["inventory_source"] == "business":
            from app.services.business_gateway import configured_gateway
            inventory_evidence = configured_gateway().inventory(data["mine_id"])
        else:
            from app.services.mysql_bridge import scoped_inventory
            inventory_evidence = scoped_inventory(data["mine_id"])
        if not inventory_evidence.get("observed_at"):
            raise ValueError("库存观测时间待配置；MySQL 接入请设置 MYSQL_INVENTORY_OBSERVED_AT")
        inventory_time = datetime.fromisoformat(inventory_evidence["observed_at"])
        age = (datetime.now(timezone.utc)-inventory_time).total_seconds()/3600
        if age < 0 or age > (config or {}).get("max_inventory_age_hours",24):
            raise ValueError("业务库存超过允许时效")
        requested = plan["parsed"]["requirements"] if plan else [
            {**item,"group":group} for group,items in (data.get("resources") or {}).items() for item in items]
        if not requested:
            raise ValueError("请提供方案资源需求后读取业务库存")
        merged = {g:[] for g in GROUPS}
        for req in requested:
            counts = inventory_evidence["available_resources"][req["group"]]
            if req["name"] not in counts:
                raise ValueError("业务库存未返回方案所需资源："+req["name"])
            merged[req["group"]].append({"resource_id":req["name"],"name":req["name"],"unit":req["unit"],
                "required":req["required"],"available":counts[req["name"]]})
        data["resources"] = validate_resources(merged)
        inventory_asset_id = None
    if plan:
        available = data.get("resources") or {g:[] for g in GROUPS}
        # 方案声明真实需求；未列明可用量单独记录，不能把缺失库存当作真实零库存。
        merged = {g:[] for g in GROUPS}; missing = []
        for req in plan["parsed"]["requirements"]:
            match = next((x for x in available.get(req["group"],[]) if x["name"]==req["name"]),None)
            if not match:
                missing.append(req["name"]); continue
            if match["unit"] != req["unit"]:
                raise ValueError("台账与方案资源单位不一致："+req["name"])
            merged[req["group"]].append({**match,"required":req["required"]})
        if missing:
            raise ValueError("方案所需资源待补充可用量："+ "、".join(missing))
        data["resources"] = validate_resources(merged)
    level = None; triggered = []
    if config and config["rules"]:
        for rule in config["rules"]:
            value = metrics.get(rule["metric"])
            if value is None:
                raise ValueError("配置所需监测指标待补充："+rule["metric"])
            if (value >= rule["threshold"]) if rule["operator"] == ">=" else (value <= rule["threshold"]):
                triggered.append({**rule,"observed":value})
        level = max((r["level"] for r in triggered),key=LEVELS.index,default="normal")
    assessment = {"method_id": "configured-rules" if config and config["rules"] else "observations-only",
                  "profile":config, "applicability":"applicable", "metrics":metrics, "risk_level":level,
                  "risk_score":None,"alarm":level in ("orange","red") if level else None,
                  "source":data["source"] if "source" in data else "用户监测摘要",
                  "window_start":start.isoformat(),"window_end":end.isoformat(),"time_source":time_source,
                  "input_files":data.get("input_files",[]),"reasons":[], "warnings":warnings,
                  "triggered_rules":triggered,
                  "summary":f"依据配置{config['version']}评估为{level}；命中{len(triggered)}条规则" if level else "已完成实际观测统计，未选择分级参数"}
    if not config and data.get("input_type") == "raw_events" and task.method_id != "statistics-v1":
        # 保留显式研究方法，但不能混入演示颜色/风险分数。
        assessment = analyze_raw(task)
        if assessment["applicability"] != "applicable":
            raise ValueError("研究方法不适用："+ "；".join(assessment["reasons"]))
    snapshot = {"task":data,"configuration":config,"documents":documents,"plan":plan,"assessment":assessment,
                "execution_mode":"actual","warnings":warnings,"inventory_evidence":inventory_evidence,
                "inventory_asset_id":inventory_asset_id}
    snapshot["snapshot_id"] = digest(snapshot)
    return snapshot
